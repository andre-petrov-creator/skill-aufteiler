#!/usr/bin/env python3
"""Idempotente Template-Migration: Bewirtschaftung (C8-Formel) + Ruecklagen-Topf-Tabelle
im Reiter "Verkauf Cashflow" der Aufteiler-Kalkulation.

Plan: plans/2026-06-26-modul6-bewirtschaftung-ruecklage.md

Setzt:
  - C8 von Wert 0,20 auf Formel (Verwaltung + Ruecklage) / Nettokaltmiete (nur Vermieterkosten)
  - Sektion 1 (Z.61-70): Bewirtschaftungs-Komponenten
  - Sektion 2 (Z.73-87): Ruecklagen-Topf-Sensitivitaet 0-50k

Re-run = gleiches Ergebnis (ueberschreibt, kein Append). Laeuft nur auf dem TEMPLATE.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

TEMPLATE = r"C:\meine-projekte\Immobilien\Aufteiler\template\Kalkulation_Aufteiler_mit_VK_CF.xlsx"

FONT = "Cambria"
C_HEADER = "FF1F3864"   # dunkelblau, Sektions-Ueberschrift
C_LABEL = "FF404040"    # dunkelgrau, Label
C_INPUT = "FF0000FF"    # blau, Input
C_FORMULA = "FF000000"  # schwarz, Formel
C_WHITE = "FFFFFFFF"
FILL_INPUT = PatternFill("solid", fgColor="FFFFF2CC")   # hellgelb (Input)
FILL_THEAD = PatternFill("solid", fgColor="FF1F3864")   # dunkelblau (Tabellenkopf)
CENTER = Alignment(horizontal="center")


def set_header(ws, coord, text):
    c = ws[coord]
    c.value = text
    c.font = Font(name=FONT, size=11, bold=True, color=C_HEADER)


def set_label(ws, coord, text):
    c = ws[coord]
    c.value = text
    c.font = Font(name=FONT, size=10, bold=False, color=C_LABEL)


def set_input(ws, coord, value, numfmt="0"):
    c = ws[coord]
    c.value = value
    c.font = Font(name=FONT, size=11, bold=False, color=C_INPUT)
    c.fill = FILL_INPUT
    c.alignment = CENTER
    c.number_format = numfmt


def set_formula(ws, coord, value, numfmt="General"):
    c = ws[coord]
    c.value = value
    c.font = Font(name=FONT, size=11, bold=False, color=C_FORMULA)
    c.fill = PatternFill(fill_type=None)
    c.alignment = CENTER
    c.number_format = numfmt


def set_thead(ws, coord, text):
    c = ws[coord]
    c.value = text
    c.font = Font(name=FONT, size=10, bold=True, color=C_WHITE)
    c.fill = FILL_THEAD
    c.alignment = CENTER


def main():
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb["Verkauf Cashflow"]

    # --- C8: Input-Wert 0,20 -> Formel (nur Vermieterkosten) ---
    set_formula(ws, "C8", "=IFERROR((C66+C67)/C69,0.2)", "0.00%")

    # --- Sektion 1: Bewirtschaftung (Z.61-70) ---
    set_header(ws, "A61", "BEWIRTSCHAFTUNG (nicht-umlagefaehig) -> speist C8")
    set_label(ws, "A62", "Verwaltung (EUR/WE/Monat)")
    set_input(ws, "C62", 35, "#,##0")
    set_label(ws, "A63", "Anzahl WE")
    set_formula(ws, "C63", "=COUNT(F18:F37)", "0")
    set_label(ws, "A64", "Laufende Ruecklage (EUR/m2/Jahr)")
    set_input(ws, "C64", 14, "#,##0.0")
    set_label(ws, "A65", "Wohnflaeche gesamt (m2)")
    set_formula(ws, "C65", "=F38", "#,##0.0")
    set_label(ws, "A66", "Verwaltung gesamt (EUR/Jahr)")
    set_formula(ws, "C66", "=C62*C63*12", "#,##0")
    set_label(ws, "A67", "Ruecklage gesamt (EUR/Jahr)")
    set_formula(ws, "C67", "=C64*C65", "#,##0")
    set_label(ws, "A68", "Ruecklage ueber 10 Jahre (EUR)")
    set_formula(ws, "C68", "=C67*10", "#,##0")
    set_label(ws, "A69", "Nettokaltmiete (EUR/Jahr)")
    set_formula(ws, "C69", "=G38*12", "#,##0")
    set_label(ws, "A70", "-> Bewirtschaftungsquote C8")
    set_formula(ws, "C70", "=C8", "0.00%")

    # --- Sektion 2: Ruecklagen-Topf-Tabelle (Z.73-87) ---
    set_header(ws, "A73", "RUECKLAGEN-TOPF -> Quote & max. Verkaufspreis")
    set_label(ws, "A74", "Vorhandene WEG-Ruecklage (Referenz)")
    set_formula(ws, "C74", "=IFERROR(BESICHTIGUNG!B11,0)", "#,##0")
    set_label(ws, "A75", "10-Jahres-Bedarf (EUR)")
    set_formula(ws, "C75", "=C68", "#,##0")

    set_thead(ws, "A76", "Topf EUR")
    set_thead(ws, "B76", "Rueckl. EUR/m2")
    set_thead(ws, "C76", "Rueckl. EUR/Jahr")
    set_thead(ws, "D76", "Quote C8")
    set_thead(ws, "E76", "max. VK-Preis")

    for i, topf in enumerate(range(0, 50001, 5000)):
        r = 77 + i
        set_formula(ws, f"A{r}", topf, "#,##0")
        set_formula(ws, f"B{r}", f"=MAX(0,($C$75-A{r})/10)/$C$65", "#,##0.0")
        set_formula(ws, f"C{r}", f"=MAX(0,($C$75-A{r})/10)", "#,##0")
        set_formula(ws, f"D{r}", f"=($C$66+C{r})/$C$69", "0.0%")
        set_formula(ws, f"E{r}", f"=(((1-$C$7)*(1-D{r})*$G$38)-$C$63*$C$11)/$C$14", "#,##0")

    wb.save(TEMPLATE)
    print("OK - Template gespeichert:", TEMPLATE)


if __name__ == "__main__":
    main()

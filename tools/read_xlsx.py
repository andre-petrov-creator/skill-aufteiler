"""XLSX-Reader für den Aufteiler-Skill (Refactor v2.0).

Liest eine Excel-Datei (typisch: Mieterliste) und gibt einen lesbaren JSON-Dump
auf stdout. Pro Sheet eine Liste mit Header + Datenzeilen + Zellen-Adressen,
sodass das LLM beim Extrahieren genau weiß: "WE-Nr in Spalte B Zeile 8" etc.

Lauf:
  python tools/read_xlsx.py <pfad-zur-xlsx>

Returncode:
  0 = OK
  1 = Datei nicht lesbar / kein XLSX
  2 = Usage-Fehler

Abhängigkeiten: openpyxl (üblich Standard in NRW-Aufteiler-Repo).
"""
import json
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError as e:
    print(f"openpyxl nicht installiert: {e}. Bitte 'pip install openpyxl'.", file=sys.stderr)
    sys.exit(1)


def _cell_value(v: Any) -> Any:
    """Excel-Zellen-Wert in JSON-taugliche Form bringen."""
    if v is None:
        return None
    # datetime → ISO-String
    if hasattr(v, "isoformat"):
        return v.isoformat()
    if isinstance(v, (int, float, bool, str)):
        return v
    return str(v)


def read_workbook(path: Path) -> dict:
    """Workbook lesen, alle Sheets als dict zurückgeben."""
    wb = load_workbook(filename=str(path), data_only=True, read_only=True)
    result: dict[str, Any] = {
        "datei": path.name,
        "sheets": {},
    }
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_dump: dict[str, Any] = {
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "zeilen": [],
        }

        for row_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):
            row_entry: dict[str, Any] = {"zeile": row_idx, "zellen": {}}
            row_empty = True
            for cell in row:
                v = _cell_value(cell.value)
                if v is None or v == "":
                    continue
                col_letter = get_column_letter(cell.column)
                row_entry["zellen"][col_letter] = v
                row_empty = False
            if not row_empty:
                sheet_dump["zeilen"].append(row_entry)

        result["sheets"][sheet_name] = sheet_dump

    wb.close()
    return result


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print("Usage: read_xlsx.py <pfad-zur-xlsx>", file=sys.stderr)
        return 2

    path = Path(argv[1])
    if not path.is_file():
        print(f"Datei nicht gefunden: {path}", file=sys.stderr)
        return 1
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        print(f"Kein XLSX/XLSM-File (Endung {path.suffix}): {path}", file=sys.stderr)
        return 1

    try:
        result = read_workbook(path)
    except Exception as e:
        print(f"Lese-Fehler: {type(e).__name__}: {e}", file=sys.stderr)
        return 1

    json.dump(result, sys.stdout, ensure_ascii=False, indent=2, default=str)
    sys.stdout.write("\n")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))

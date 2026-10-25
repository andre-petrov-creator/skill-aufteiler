#!/usr/bin/env python3
"""Aufteiler Modul 6: state.json -> Kalkulations-Excel (Template-Kopie befuellen).

Lauf:   python tools/export_excel.py "runs/<slug>/state.json" [--out PFAD] [--dry-run] [--force]

Portiert den Excel-Write-Block aus skills/aufteiler-modul-6-deal-bewertung/SKILL.md
(Sektion 3b) in ein reproduzierbares Werkzeug. Zell-Vertraege: docs/excel_handoff.md.

--dry-run  druckt nur die Werte-Vorschau (Sheet!Zelle | Wert | Quelle) und schreibt nichts -
           das ist die Pflicht-Vorschau aus dem Modul-6-Skill vor jeder Excel-Freigabe.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment

sys.path.insert(0, str(Path(__file__).resolve().parent))

REPO_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = REPO_ROOT / "template" / "Kalkulation_Aufteiler_mit_VK_CF.xlsx"
COMMENT_AUTHOR = "Aufteiler-Skill"
REACH_TIME_DEFAULT_MO = 60
RND_GUTACHTEN_NETTO_JE_WE = 600
WEG_TEILUNG_TEMPLATE_DEFAULT = 12000
NRW_KAPPUNGSGRENZE = 0.15

# Modul-4-Kategorie -> (RENO-Zeile, Mengen-Schluessel); Details docs/excel_handoff.md
KATEGORIE_RENO_MAPPING = {
    "Dach": (25, "dachflaeche"),
    "Fassade": (18, "fassadenflaeche"),
    "Fenster": (50, "fenster_stk"),
    "Heizung": (78, "anzahl_we"),
    "Elektrik": (69, "wohnflaeche_gesamt"),
    "Sanitaer": (93, "wohnflaeche_gesamt"),
    "Boeden": (37, "wohnflaeche_gesamt"),
    "Grundriss": (30, "anzahl_we"),
}


class Writes:
    """Sammelt alle geplanten Zell-Writes, damit Vorschau und Schreiben dieselbe Liste nutzen."""

    def __init__(self) -> None:
        self.items: list[tuple[str, str, object, str]] = []
        self.comments: list[tuple[str, str, str]] = []

    def set(self, sheet: str, cell: str, value, source: str) -> None:
        if value is None:
            return
        self.items.append((sheet, cell, value, source))

    def comment(self, sheet: str, cell: str, text: str) -> None:
        self.comments.append((sheet, cell, text))

    def apply(self, wb) -> None:
        for sheet, cell, value, _ in self.items:
            wb[sheet][cell] = value
        for sheet, cell, text in self.comments:
            wb[sheet][cell].comment = Comment(text, COMMENT_AUTHOR)

    def preview(self) -> str:
        width = max((len(f"{s}!{c}") for s, c, _, _ in self.items), default=10)
        lines = [f"{'Zelle':<{width}} | Wert | Quelle", "-" * (width + 40)]
        for sheet, cell, value, source in self.items:
            shown = value.isoformat() if isinstance(value, dt.date) else value
            lines.append(f"{sheet + '!' + cell:<{width}} | {shown} | {source}")
        lines.append(f"{len(self.items)} Zellen, {len(self.comments)} Kommentare")
        return "\n".join(lines)


def strasse_kurz(slug: str) -> str:
    if ", " in slug:
        return slug.split(", ")[0].replace(" ", "_")
    return "-".join(slug.split("-")[:2])


def as_date(value) -> dt.date | None:
    """ISO-String -> date. Excel braucht echte Datumswerte, sonst schlaegt ISNUMBER fehl."""
    if not value:
        return None
    if isinstance(value, dt.date):
        return value
    return dt.date.fromisoformat(str(value)[:10])


def etage_aus_lage(lage: str) -> str:
    L = lage.lower()
    if L.startswith("eg"):
        return "EG"
    if "dg" in L or "dachgeschoss" in L:
        return "DG"
    if "3.og" in L or "3_og" in L or "og3" in L:
        return "3.OG"
    if "2.og" in L or "2_og" in L or "og2" in L:
        return "2.OG"
    if "og" in L:
        return "1.OG"
    return lage


def lage_aus_lage(lage: str) -> str:
    L = lage.lower()
    if "links" in L:
        return "links"
    if "rechts" in L:
        return "rechts"
    return ""


def ruecklage_eur_qm(rnd_jahre: float) -> int:
    return 10 if rnd_jahre >= 40 else (14 if rnd_jahre >= 20 else 18)


def load_objektdaten(state_path: Path) -> dict:
    p = state_path.parent / "objektdaten.json"
    if not p.is_file():
        return {}
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}


def wert(node):
    """objektdaten.json fuehrt jeden Wert als {wert, quelle, zitat, konfidenz}."""
    return node.get("wert") if isinstance(node, dict) else node


def ist_miete_monat(miete: dict, we: dict, od: dict) -> float:
    """Sollstellungs-Betrag in EUR hat Vorrang; qm x Flaeche verliert Cents durch das gerundete EUR/m2."""
    for candidate in (miete.get("ist_miete_eur_pro_monat"), wert(od.get("ist_miete_eur_monat"))):
        if candidate is not None:
            return round(float(candidate), 2)
    return round(miete["ist_miete_eur_pro_qm"] * we["wohnflaeche_qm"], 2)


def template_is_v2(template: Path) -> bool:
    """Migriertes Template traegt den Defined Name AUFTEILER_TPL_V2 (Reiter Verkauf Cashflow als Vorwaertskalkulation)."""
    wb = load_workbook(template, read_only=True)
    try:
        return "AUFTEILER_TPL_V2" in wb.defined_names
    finally:
        wb.close()


def plan_writes(state: dict, objektdaten: dict, is_v2: bool = False) -> Writes:
    w = Writes()
    m1 = state.get("modul_1") or {}
    m2 = state["modul_2"]
    m3 = state["modul_3"]
    m4 = state["modul_4"]
    m5 = state["modul_5"]
    m6 = state.get("modul_6") or {}
    we_liste = m2["we_liste"]
    we_mieten = m5["we_mieten"]
    n_we = len(we_liste)
    wohnflaeche_gesamt = sum(we["wohnflaeche_qm"] for we in we_liste)
    od_we = {wert(x.get("we_nr")): x for x in objektdaten.get("we_liste", [])}

    # ----- 1. BESICHTIGUNG -----
    w.set("BESICHTIGUNG", "B6", state["objekt"]["adresse"], "objekt.adresse")
    w.set("BESICHTIGUNG", "B7", wohnflaeche_gesamt, "Summe modul_2.we_liste[].wohnflaeche_qm")
    w.set("BESICHTIGUNG", "B8", m3["baujahr"], "modul_3.baujahr")
    w.set("BESICHTIGUNG", "B9", n_we, "len(modul_2.we_liste)")
    w.set("BESICHTIGUNG", "B13", m1.get("angebotspreis_eur"), "modul_1.angebotspreis_eur")
    ist_monat = [ist_miete_monat(m, we, od_we.get(we["we_nr"], {})) for m, we in zip(we_mieten, we_liste)]
    w.set("BESICHTIGUNG", "B33", round(sum(ist_monat) * 12, 2), "Summe Ist-Kaltmiete x 12 (Modul 5)")
    kappung = m5.get("kappungsgrenze", NRW_KAPPUNGSGRENZE)
    w.set("BESICHTIGUNG", "B34", kappung, "modul_5.kappungsgrenze (Default NRW 0,15)")

    # ----- 2. MIETER -----
    for i, (we, miete) in enumerate(zip(we_liste, we_mieten)):
        row = 8 + i
        od = od_we.get(we["we_nr"], {})
        w.set("MIETER", f"A{row}", 1, "Haus-Nr (Single-Haus)")
        w.set("MIETER", f"B{row}", we["we_nr"], "modul_2.we_liste[].we_nr")
        w.set("MIETER", f"C{row}", etage_aus_lage(we["lage_im_haus"]), "abgeleitet aus lage_im_haus")
        w.set("MIETER", f"D{row}", lage_aus_lage(we["lage_im_haus"]), "abgeleitet aus lage_im_haus")
        w.set("MIETER", f"E{row}", we["zimmer_anzahl"], "modul_2.we_liste[].zimmer_anzahl")
        w.set("MIETER", f"F{row}", we["wohnflaeche_qm"], "modul_2.we_liste[].wohnflaeche_qm")
        w.set("MIETER", f"G{row}", wert(od.get("mieter_name")), "objektdaten.json we_liste[].mieter_name")
        w.set("MIETER", f"H{row}", as_date(miete.get("mieter_seit") or wert(od.get("mieter_seit"))),
              "modul_5.we_mieten[].mieter_seit (Fallback objektdaten.json)")
        w.set("MIETER", f"I{row}", as_date(miete.get("letzte_mieterhoehung")), "modul_5.we_mieten[].letzte_mieterhoehung")
        w.set("MIETER", f"J{row}", ist_monat[i], "modul_5.we_mieten[].ist_miete_eur_pro_monat (Fallback qm x Flaeche)")
        # MIETER!K ist die Stellplatz-/Garagenmiete in EUR (K28 = SUM), keine Stueckzahl.
        # Im migrierten Template (V2) ist K ein Formel-Spiegel aus Verkauf Cashflow AL+AM - dann nicht schreiben.
        if not is_v2:
            asset_miete = (miete.get("garage_miete_eur_pro_monat") or 0) + (miete.get("stellplatz_miete_eur_pro_monat") or 0)
            w.set("MIETER", f"K{row}", asset_miete if asset_miete else None, "garage_miete + stellplatz_miete (Modul 5)")
        w.set("MIETER", f"Y{row}", miete["mietspiegel_obergrenze_eur_pro_qm"], "modul_5.we_mieten[].mietspiegel_obergrenze_eur_pro_qm")
    m6_mittel = sum(m["sollmiete_eur_pro_qm"] * we["wohnflaeche_qm"] for m, we in zip(we_mieten, we_liste)) / wohnflaeche_gesamt
    w.set("MIETER", "M6", round(m6_mittel, 2), "gewichteter Mietspiegel-Mittelwert (Modul 5)")
    w.set("MIETER", "P6", kappung, "Kappungsgrenze (Modul 5)")
    w.set("MIETER", "T4", m5.get("paragraph_559_umlage_prozent"), "modul_5.paragraph_559_umlage_prozent")

    # ----- 3. RENO -----
    weg = m4.get("weg_teilung_netto_eur") or 0
    if weg > 0:
        w.set("RENO", "I104", 1, "WEG-Teilung Menge")
        if abs(weg - WEG_TEILUNG_TEMPLATE_DEFAULT) > 100:
            w.set("RENO", "K104", weg, "modul_4.weg_teilung_netto_eur (weicht vom Default ab)")
    w.set("RENO", "F105", "RND", "RND-Gutachten Bezeichnung")
    w.set("RENO", "I105", n_we, "RND-Gutachten Menge = Anzahl WE")
    w.set("RENO", "K105", RND_GUTACHTEN_NETTO_JE_WE, "RND-Gutachten 600 EUR netto/WE")
    mengen = {
        "wohnflaeche_gesamt": wohnflaeche_gesamt,
        "dachflaeche": wohnflaeche_gesamt * 0.6,
        "fassadenflaeche": wohnflaeche_gesamt * 0.8,
        "fenster_stk": 6 * n_we,
        "anzahl_we": n_we,
    }
    for m in m4.get("massnahmen_liste", []):
        hit = KATEGORIE_RENO_MAPPING.get(m["kategorie"])
        if not hit:
            continue
        zeile, key = hit
        w.set("RENO", f"I{zeile}", round(mengen[key], 0), f"Menge {m['kategorie']} (Heuristik)")

    # ----- Mietsubvention -> VERKAUFSMATRIX!Y27 (negativ) -----
    total_subv = m5.get("mietsubventionen_total_eur", m5["mietsubventionen_summe_eur_pro_monat"] * REACH_TIME_DEFAULT_MO)
    w.set("VERKAUFSMATRIX", "Y27", -round(total_subv, 2), "modul_5.mietsubventionen_total_eur (negativ)")

    # ----- 4. KALKU Portalpreise -----
    eur_per_qm = m1.get("eur_per_qm") or {}
    for cell, key in (("I13", "homeday"), ("I14", "immometrica"), ("I15", "check24"), ("I16", "interhyp"), ("I17", "grundstuecksmarktbericht")):
        val = eur_per_qm.get(key)
        w.set("KALKU", cell, float(val) if val is not None else None, f"modul_1.eur_per_qm.{key}")

    # ----- 5. Verkauf Cashflow -----
    # V1-Template: feste Zellen; V2: Adressen aus dem Vertrag (Parameterblock liegt unter der Tabelle)
    def vkcf_cell(key: str, v1_cell: str) -> str:
        if not is_v2:
            return v1_cell
        import vkcf_vertrag as V
        return V.vertrag().param_cell(key)

    w.set("Verkauf Cashflow", vkcf_cell("afa_satz_std", "C9"), round(m3["afa_empfehlung_prozent"] / 100, 4), "modul_3.afa_empfehlung_prozent / 100")
    w.set("Verkauf Cashflow", vkcf_cell("gebaeude_anteil_std", "C10"), round(m2["gebaeude_anteil_prozent"] / 100, 4), "modul_2.gebaeude_anteil_prozent / 100")
    w.set("Verkauf Cashflow", vkcf_cell("ruecklage_eur_qm_jahr", "C64"), ruecklage_eur_qm(m3["rnd_jahre"]), "RND-Staffel 10/14/18 aus modul_3.rnd_jahre")

    # ----- 5b. Verkauf Cashflow: Datenblock je WE (nur im migrierten Template V2) -----
    if is_v2:
        plan_vkcf_block(w, state, m1, m2, we_liste, we_mieten)

    # ----- 6. Kommentare -----
    korridor = m3.get("afa_korridor_prozent") or {}
    w.comment("KALKU", "H10",
              f"Aufteiler-Skill Modul 3 RND/AfA:\n"
              f"AfA-Empfehlung: {m3['afa_empfehlung_prozent']:.2f}% (Korridor "
              f"{korridor.get('min', 0):.2f}-{korridor.get('max', 0):.2f}%)\n"
              f"RND: {m3['rnd_jahre']} Jahre (rnd_frozen=true)\n"
              f"Mod-Score: {m3.get('mod_score', '-')}/20\n"
              f"BFH IX R 7/23: Gutachten zwingend, dies ist eine Vorpruefung.")
    if m6.get("bewertungs_score") is not None:
        w.comment("BESICHTIGUNG", "A40",
                  f"Aufteiler-Skill Modul 6 Deal-Bewertung:\n"
                  f"Score: {m6['bewertungs_score']}/100\nStatus: {str(m6.get('status', '')).upper()}\n"
                  f"PDF: {m6.get('pdf_pfad', '-')}")
    w.comment("VERKAUFSMATRIX", "Y27",
              f"Aufteiler-Skill Modul 5 Mietsubvention (Subventionskosten, negativ):\n"
              f"Summe EUR/Monat: {m5['mietsubventionen_summe_eur_pro_monat']:.2f}\n"
              f"Gesamt ueber Reach-Time: {total_subv:.2f} EUR\n"
              f"Stufen-Details siehe runs/<slug>/modul-5-output.md.")
    return w


def plan_vkcf_block(w: Writes, state: dict, m1: dict, m2: dict, we_liste: list, we_mieten: list) -> None:
    """Vorwaertskalkulation je Wohnung: Startwerte in die Eingabespalten des Reiters "Verkauf Cashflow".

    Adressen kommen ausschliesslich ueber den Spaltenschluessel aus tools/vkcf_vertrag.py.
    Solange dort SPALTEN leer ist (Template noch nicht migriert), passiert nichts.
    """
    import vkcf_vertrag as V  # lokal, damit das Werkzeug auch ohne Vertrag laeuft

    if not V.SPALTEN:
        return
    vt = V.vertrag()
    sheet = V.SHEET

    def cell(key: str, row: int) -> str | None:
        return f"{vt.letter[key]}{row}" if key in vt.letter else None

    def put(key: str, row: int, value, source: str) -> None:
        c = cell(key, row)
        if c is not None and value is not None:
            w.set(sheet, c, value, source)

    median = (m1.get("eur_per_qm") or {}).get("median")
    for i, (we, miete) in enumerate(zip(we_liste, we_mieten)):
        row = V.LAYOUT.first + i
        # nur Eingabespalten; alles andere (Preise, Cashflow, Subvention gesamt) rechnet das Blatt
        put("abgabe_eur_qm_regler", row, float(median) if median is not None else None, "modul_1.eur_per_qm.median (Startwert, je WE veraenderbar)")
        put("mea_zaehler", row, we.get("mea_zaehler"), "modul_2.we_liste[].mea_zaehler")
        if "leerstand" in miete:
            put("leerstand", row, "ja" if miete["leerstand"] else "nein", "modul_5.we_mieten[].leerstand")
        put("garagen_anzahl", row, miete.get("garagen_anzahl") or None, "modul_5.we_mieten[].garagen_anzahl")
        put("stellplaetze_anzahl", row, miete.get("stellplaetze_anzahl") or None, "modul_5.we_mieten[].stellplaetze_anzahl")
        # Verkaufspreis je Asset = Anzahl je WE (Modul 5) x Preis je Einheit (Modul 2, Marktansatz mit Gate)
        for anz_key, preis_key, col in (("garagen_anzahl", "garage_preis_eur", "garage_preis"),
                                        ("stellplaetze_anzahl", "stellplatz_preis_eur", "stellplatz_preis")):
            anz, preis = miete.get(anz_key) or 0, m2.get(preis_key)
            if anz and preis is not None:
                put(col, row, round(anz * float(preis), 2), f"{anz_key} x modul_2.{preis_key} ({m2.get('garagen_preis_quelle', 'Quelle offen')})")
        # Stufen-Ziel von Modul 5 (Zone A "Ziel EUR/m2") = Mittelwert aus Sollmiete (Mittelwert + Zuschlaege) und Obergrenze;
        # darauf rechnen reach_time_monate und mietsubvention_eur_pro_monat - nicht auf sollmiete_eur_pro_qm allein
        ziel_qm = (miete["sollmiete_eur_pro_qm"] + miete["mietspiegel_obergrenze_eur_pro_qm"]) / 2
        put("soll_miete", row, round(ziel_qm * we["wohnflaeche_qm"], 2), "Ziel-Sollmiete = (sollmiete + obergrenze)/2 x wohnflaeche_qm (Modul 5 Stufen-Ziel)")
        put("mietsubvention_eur_pro_monat", row, miete.get("mietsubvention_eur_pro_monat"), "modul_5.we_mieten[].mietsubvention_eur_pro_monat (Stufendurchschnitt)")
        put("garage_miete", row, miete.get("garage_miete_eur_pro_monat") or None, "modul_5.we_mieten[].garage_miete_eur_pro_monat")
        put("stellplatz_miete", row, miete.get("stellplatz_miete_eur_pro_monat") or None, "modul_5.we_mieten[].stellplatz_miete_eur_pro_monat")
        put("reach_time_monate", row, miete.get("reach_time_monate"), "modul_5.we_mieten[].reach_time_monate")
        put("bewirtschaftungsquote_schaetzung", row, miete.get("bewirtschaftungsquote_schaetzung"), "modul_5.we_mieten[].bewirtschaftungsquote_schaetzung")
        put("hausgeld_nicht_umlagefaehig", row, miete.get("hausgeld_nicht_umlagefaehig_eur_pro_monat"), "modul_5.we_mieten[].hausgeld_nicht_umlagefaehig_eur_pro_monat")
        put("hausgeld_ruecklage", row, miete.get("hausgeld_ruecklage_eur_pro_monat"), "modul_5.we_mieten[].hausgeld_ruecklage_eur_pro_monat")
    # Parameterblock: Grundstueck, BRW, MEA-Nenner aus Modul 2 (Zellen laut Vertrag, Quelle als Kommentar)
    gq = m2.get("grundstuecksflaeche_quelle") or {}
    gq_text = ", ".join(str(gq[k]) for k in ("quelle", "dokument", "flurstueck", "konfidenz") if gq.get(k)) or None
    c_grund, c_brw, c_mea = vt.param_cell("grundstueck_qm"), vt.param_cell("brw"), vt.param_cell("mea_nenner")
    w.set(sheet, c_grund, m2.get("grundstuecksflaeche_qm"), "modul_2.grundstuecksflaeche_qm")
    if gq_text and m2.get("grundstuecksflaeche_qm") is not None:
        w.comment(sheet, c_grund, f"Quelle: {gq_text}")
    w.set(sheet, c_brw, m2.get("brw_eur_pro_qm"), "modul_2.brw_eur_pro_qm")
    w.set(sheet, c_mea, m2.get("mea_nenner"), "modul_2.mea_nenner")
    if m2.get("mea_quelle") and m2.get("mea_nenner") is not None:
        w.comment(sheet, c_mea, f"MEA-Quelle: {m2['mea_quelle']}")


def write_workbook(writes: Writes, out: Path, template: Path, force: bool, is_v2: bool = False) -> None:
    if out.exists() and not force:
        raise SystemExit(f"Zieldatei existiert bereits: {out}\n--force ueberschreibt (Backup wird angelegt).")
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists():
        backup = out.with_name(f"{out.stem}.bak-{dt.datetime.now():%Y%m%d-%H%M%S}{out.suffix}")
        shutil.copy2(out, backup)
        print(f"Backup: {backup}")
    # Nie das Ziel direkt oeffnen: Temp-Datei fuellen, dann atomar ersetzen
    fd, tmp_name = tempfile.mkstemp(suffix=".xlsx", dir=out.parent)
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        shutil.copyfile(template, tmp)
        wb = load_workbook(tmp)
        writes.apply(wb)
        wb.save(tmp)
        if is_v2:
            # openpyxl drops sheetFormatPr.outlineLevelCol on save -> Excel hides the +/- buttons of the collapsed helper columns
            from migrate_vkcf_vorwaertskalkulation import patch_outline_level
            patch_outline_level(tmp, "Verkauf Cashflow")
        os.replace(tmp, out)
    finally:
        if tmp.exists():
            tmp.unlink()


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("state", help="Pfad zu runs/<slug>/state.json")
    ap.add_argument("--out", help="Ziel-Excel (Default: runs/<slug>/Kalkulation_<Strasse>.xlsx)")
    ap.add_argument("--template", default=str(TEMPLATE))
    ap.add_argument("--dry-run", action="store_true", help="nur Werte-Vorschau, nichts schreiben")
    ap.add_argument("--force", action="store_true", help="bestehende Zieldatei ueberschreiben (mit Backup)")
    args = ap.parse_args()

    state_path = Path(args.state)
    state = json.loads(state_path.read_text(encoding="utf-8"))
    is_v2 = template_is_v2(Path(args.template))
    print(f"Template: {'V2 (Vorwaertskalkulation je WE)' if is_v2 else 'V1 (Rueckrechnung) - Datenblock wird nicht geschrieben'}")
    writes = plan_writes(state, load_objektdaten(state_path), is_v2)
    print(writes.preview())
    if args.dry_run:
        return 0
    out = Path(args.out) if args.out else state_path.parent / f"Kalkulation_{strasse_kurz(state['objekt']['slug'])}.xlsx"
    write_workbook(writes, out, Path(args.template), args.force, is_v2)
    print(f"Geschrieben: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
